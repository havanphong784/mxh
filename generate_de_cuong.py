#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tạo file Đề cương Đồ án Cơ sở 2 chuẩn biểu mẫu VKU KHMT
Đề tài: Xây dựng website mạng xã hội chia sẻ bài viết và tương tác trực tuyến
Dành cho nhóm 02 sinh viên - Tinh gọn, đúng tầm Đồ án cơ sở 2, chuẩn 2 trang A4.
"""

import os
import shutil
import zipfile
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_font(run, name="Times New Roman", size_pt=10.0, bold=False, italic=False, color_rgb=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    rFonts.set(qn('w:eastAsia'), name)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for edge, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, border_color: str | None = "CBD5E1", border_sz: int = 4, inside: bool = True):
    tblPr = table._tbl.tblPr
    tblBorders = tblPr.find(qn('w:tblBorders'))
    if tblBorders is not None:
        tblPr.remove(tblBorders)
    
    if border_color is None:
        new_borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
                <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
                <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
                <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
                <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
                <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            </w:tblBorders>
        ''')
    else:
        inside_xml = f'''
            <w:insideH w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
        ''' if inside else '''
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        '''
        new_borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
                <w:left w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
                <w:bottom w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
                <w:right w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>
                {inside_xml}
            </w:tblBorders>
        ''')
    tblPr.append(new_borders)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def create_document():
    doc = docx.Document()
    
    # Cấu hình Margins A4 chuẩn (Top: 1.6cm, Bottom: 1.6cm, Left: 2.3cm, Right: 1.6cm)
    section = doc.sections[0]
    section.page_width = Pt(595.3)   # A4: 21.0 cm
    section.page_height = Pt(841.9)  # A4: 29.7 cm
    
    section.top_margin = Pt(45.35)    # ~1.6 cm
    section.bottom_margin = Pt(45.35) # ~1.6 cm
    section.left_margin = Pt(65.0)    # ~2.3 cm
    section.right_margin = Pt(45.35)  # ~1.6 cm
    
    content_width_pt = 595.3 - 65.0 - 45.35 # 484.95 pt
    
    # ==========================================
    # HEADER: Bảng Table 0 (Không viền, 2 cột)
    # ==========================================
    table0 = doc.add_table(rows=1, cols=2)
    table0.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table0, border_color=None)
    
    col_w_half = Pt(content_width_pt / 2.0)
    for row in table0.rows:
        row.cells[0].width = col_w_half
        row.cells[1].width = col_w_half
        set_cell_margins(row.cells[0], top=15, bottom=15, left=20, right=20)
        set_cell_margins(row.cells[1], top=15, bottom=15, left=20, right=20)
    
    p0_left = table0.rows[0].cells[0].paragraphs[0]
    p0_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0_left.paragraph_format.line_spacing = 1.15
    p0_left.paragraph_format.space_after = Pt(0)
    
    r = p0_left.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN\nVÀ TRUYỀN THÔNG VIỆT - HÀN\n")
    set_font(r, "Times New Roman", size_pt=9.5, bold=True)
    r = p0_left.add_run("KHOA KHOA HỌC MÁY TÍNH\n")
    set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p0_left.add_run("------------------------")
    set_font(r, "Times New Roman", size_pt=9.0, bold=False, color_rgb=RGBColor(0x78, 0x78, 0x78))
    
    p0_right = table0.rows[0].cells[1].paragraphs[0]
    p0_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0_right.paragraph_format.line_spacing = 1.15
    p0_right.paragraph_format.space_after = Pt(0)
    
    r = p0_right.add_run("CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    set_font(r, "Times New Roman", size_pt=9.5, bold=True)
    r = p0_right.add_run("Độc lập - Tự do - Hạnh phúc\n")
    set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p0_right.add_run("------------------------------------\n")
    set_font(r, "Times New Roman", size_pt=9.0, bold=False, color_rgb=RGBColor(0x78, 0x78, 0x78))
    r = p0_right.add_run("Đà Nẵng, ngày 08 tháng 10 năm 2026")
    set_font(r, "Times New Roman", size_pt=9.5, bold=False, italic=True)
    
    # ==========================================
    # TIÊU ĐỀ CHÍNH
    # ==========================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(5)
    r = p_title.add_run("ĐỀ CƯƠNG ĐỒ ÁN CƠ SỞ 2")
    set_font(r, "Times New Roman", size_pt=14.0, bold=True)
    
    # ==========================================
    # 1. THÔNG TIN NGƯỜI THỰC HIỆN
    # ==========================================
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(2)
    p_sec1.paragraph_format.space_after = Pt(3)
    r = p_sec1.add_run("1. Thông tin người thực hiện:")
    set_font(r, "Times New Roman", size_pt=11.0, bold=True)
    
    # Bảng 2 rows x 2 cols, viền thanh lịch
    table1 = doc.add_table(rows=2, cols=2)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table1, border_color="CBD5E1", border_sz=4, inside=True)
    
    table1.rows[0].cells[0].width = Pt(content_width_pt * 0.52)
    table1.rows[0].cells[1].width = Pt(content_width_pt * 0.48)
    table1.rows[1].cells[0].width = Pt(content_width_pt * 0.52)
    table1.rows[1].cells[1].width = Pt(content_width_pt * 0.48)
    
    for row in table1.rows:
        for c in row.cells:
            set_cell_margins(c, top=50, bottom=50, left=70, right=70)
    
    # Row 0, Col 0: Sinh viên 1
    c00 = table1.rows[0].cells[0]
    p = c00.paragraphs[0]
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Họ và tên SV 1: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("Hà Văn Phong\n"); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("MSSV: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("25ITB126           "); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    r = p.add_run("Lớp: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("25IT3\n"); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    r = p.add_run("Ngành: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("Công nghệ thông tin\n"); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    r = p.add_run("Email: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("havanphong784@gmail.com\n"); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    r = p.add_run("Phụ trách: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("Cơ sở dữ liệu, Lập trình Backend & API"); set_font(r, "Times New Roman", size_pt=9.5, bold=False)
    
    # Row 0, Col 1: Sinh viên 2
    c01 = table1.rows[0].cells[1]
    p = c01.paragraphs[0]
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Họ và tên SV 2: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("<Họ và tên SV 2>\n"); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("MSSV: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("<Mã số SV 2>    "); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    r = p.add_run("Lớp: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("25IT3\n"); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    r = p.add_run("Ngành: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("Công nghệ thông tin\n"); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    r = p.add_run("Email: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("<email_sv2@vku.udn.vn>\n"); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    r = p.add_run("Phụ trách: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("Thiết kế giao diện Frontend & UI/UX"); set_font(r, "Times New Roman", size_pt=9.5, bold=False)
    
    # Row 1, Col 0: GVHD
    c10 = table1.rows[1].cells[0]
    p = c10.paragraphs[0]
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Giảng viên hướng dẫn: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("TS. Nguyễn Đức Hiển"); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    
    # Row 1, Col 1: Hình thức
    c11 = table1.rows[1].cells[1]
    p = c11.paragraphs[0]
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Hình thức thực hiện: "); set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p.add_run("Nhóm 02 sinh viên"); set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    
    # ==========================================
    # 2. TÊN ĐỒ ÁN
    # ==========================================
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(5)
    p_sec2.paragraph_format.space_after = Pt(3)
    p_sec2.paragraph_format.line_spacing = 1.2
    r = p_sec2.add_run("2. Tên đồ án: ")
    set_font(r, "Times New Roman", size_pt=11.0, bold=True)
    r = p_sec2.add_run("XÂY DỰNG WEBSITE MẠNG XÃ HỘI CHIA SẺ BÀI VIẾT VÀ TƯƠNG TÁC TRỰC TUYẾN")
    set_font(r, "Times New Roman", size_pt=10.5, bold=True)
    
    # ==========================================
    # 3. MÔ TẢ
    # ==========================================
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(3)
    p_sec3.paragraph_format.space_after = Pt(2)
    p_sec3.paragraph_format.line_spacing = 1.2
    r = p_sec3.add_run("3. Mô tả: ")
    set_font(r, "Times New Roman", size_pt=11.0, bold=True)
    r = p_sec3.add_run("(mô tả tổng quan đề tài, mục tiêu dự kiến sẽ đạt được, kết quả đề tài là gì)")
    set_font(r, "Times New Roman", size_pt=9.5, bold=False, italic=True)
    
    # - Mô tả tổng quan đề tài
    p_mota = doc.add_paragraph()
    p_mota.paragraph_format.space_before = Pt(1)
    p_mota.paragraph_format.space_after = Pt(2)
    p_mota.paragraph_format.line_spacing = 1.2
    r = p_mota.add_run("- Mô tả tổng quan đề tài: ")
    set_font(r, "Times New Roman", size_pt=10.5, bold=True)
    r = p_mota.add_run(
        "Xây dựng ứng dụng web mạng xã hội hoàn chỉnh theo mô hình Client - Server hiện đại. Hệ thống cung cấp các chức năng "
        "cốt lõi cho người dùng bao gồm: đăng ký và xác thực tài khoản qua email, tạo trang hồ sơ cá nhân, đăng tải bài viết "
        "kèm hình ảnh, tương tác thích (like), bình luận thảo luận theo chuỗi và theo dõi (follow) người dùng khác."
    )
    set_font(r, "Times New Roman", size_pt=10.0, bold=False)
    
    # - Tech Stack
    p_tech_title = doc.add_paragraph()
    p_tech_title.paragraph_format.space_before = Pt(2)
    p_tech_title.paragraph_format.space_after = Pt(1.5)
    p_tech_title.paragraph_format.line_spacing = 1.2
    r = p_tech_title.add_run("- Tech Stack sử dụng trong đề tài:")
    set_font(r, "Times New Roman", size_pt=10.5, bold=True)
    
    tech_bullets = [
        ("Frontend: ", "Next.js, React, TypeScript, Tailwind CSS, Zustand."),
        ("Backend & CSDL: ", "NestJS, Fastify, PostgreSQL, Prisma ORM."),
        ("Lưu trữ Media: ", "Cloudinary API (quản lý tải lên và lưu trữ hình ảnh/video)."),
        ("Xác thực & Email: ", "JWT (JSON Web Token), Mã hóa mật khẩu Argon2, Gửi mã OTP xác thực qua Nodemailer.")
    ]
    for label, desc in tech_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        p_b.paragraph_format.line_spacing = 1.18
        p_b.paragraph_format.left_indent = Pt(12)
        r = p_b.add_run("● ")
        set_font(r, "Times New Roman", size_pt=8.0, bold=False)
        r = p_b.add_run(label)
        set_font(r, "Times New Roman", size_pt=9.5, bold=True)
        r = p_b.add_run(desc)
        set_font(r, "Times New Roman", size_pt=9.5, bold=False)
    
    # - Mục tiêu dự kiến
    p_obj_title = doc.add_paragraph()
    p_obj_title.paragraph_format.space_before = Pt(2)
    p_obj_title.paragraph_format.space_after = Pt(1.5)
    p_obj_title.paragraph_format.line_spacing = 1.2
    r = p_obj_title.add_run("- Mục tiêu dự kiến:")
    set_font(r, "Times New Roman", size_pt=10.5, bold=True)
    
    obj_bullets = [
        "Xây dựng giao diện web mạng xã hội hiện đại, thân thiện và tương thích tốt trên cả máy tính lẫn điện thoại (Responsive).",
        "Thiết kế Cơ sở dữ liệu quan hệ chuẩn hóa trên PostgreSQL với Prisma ORM, quản lý chặt chẽ ràng buộc khóa ngoại và quan hệ giữa các bảng.",
        "Phát triển hệ thống Backend RESTful API hoàn chỉnh đáp ứng đầy đủ các thao tác CRUD cho bài viết, người dùng, tương tác và bình luận.",
        "Tích hợp tải lên hình ảnh/video trực tiếp lên dịch vụ đám mây Cloudinary an toàn và tối ưu dung lượng hiển thị.",
        "Hoàn thiện hệ thống bảo mật cơ bản: mã hóa mật khẩu, phân quyền truy cập bằng JWT và xác thực người dùng qua OTP email."
    ]
    for desc in obj_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        p_b.paragraph_format.line_spacing = 1.18
        p_b.paragraph_format.left_indent = Pt(12)
        r = p_b.add_run("● ")
        set_font(r, "Times New Roman", size_pt=8.0, bold=False)
        r = p_b.add_run(desc)
        set_font(r, "Times New Roman", size_pt=9.5, bold=False)
    
    # - Kết quả đạt được
    p_res_title = doc.add_paragraph()
    p_res_title.paragraph_format.space_before = Pt(2)
    p_res_title.paragraph_format.space_after = Pt(1.5)
    p_res_title.paragraph_format.line_spacing = 1.2
    r = p_res_title.add_run("- Kết quả đạt được:")
    set_font(r, "Times New Roman", size_pt=10.5, bold=True)
    
    res_bullets = [
        "Hệ thống website mạng xã hội hoạt động thực tế với đầy đủ các tính năng người dùng cơ bản.",
        "Bộ mã nguồn hoàn chỉnh, tài liệu đặc tả API, báo cáo Đề cương Đồ án cơ sở 2 và slide thuyết trình bảo vệ."
    ]
    for desc in res_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        p_b.paragraph_format.line_spacing = 1.18
        p_b.paragraph_format.left_indent = Pt(12)
        r = p_b.add_run("● ")
        set_font(r, "Times New Roman", size_pt=8.0, bold=False)
        r = p_b.add_run(desc)
        set_font(r, "Times New Roman", size_pt=9.5, bold=False)
    
    # ==========================================
    # 4. NỘI DUNG THỰC HIỆN
    # ==========================================
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(3)
    p_sec4.paragraph_format.space_after = Pt(2)
    p_sec4.paragraph_format.line_spacing = 1.2
    r = p_sec4.add_run("4. Nội dung thực hiện: ")
    set_font(r, "Times New Roman", size_pt=11.0, bold=True)
    r = p_sec4.add_run("(các nội dung sẽ thực hiện trong đề tài)")
    set_font(r, "Times New Roman", size_pt=9.5, bold=False, italic=True)
    
    tasks_content = [
        ("Phân tích hệ thống: ", "Khảo sát yêu cầu bài toán, xác định các thực thể (User, Post, Comment, Like, Follow). Xây dựng sơ đồ Use Case và biểu đồ quy trình nghiệp vụ."),
        ("Thiết kế Cơ sở dữ liệu: ", "Xây dựng sơ đồ ERD, thiết kế lược đồ quan hệ các bảng trong PostgreSQL và cài đặt qua Prisma ORM."),
        ("Xây dựng Backend: ", "Lập trình các module API cho Xác thực (Auth), Quản lý người dùng (User Profile), Đăng bài viết (Post CRUD), Bình luận (Comment) và Tương tác (Like, Follow)."),
        ("Lập trình Giao diện: ", "Thiết kế các trang giao diện Web: Trang Đăng nhập/Đăng ký, Bảng tin (Newsfeed), Soạn bài viết kèm ảnh, Trang cá nhân và hiển thị bình luận."),
        ("Kiểm thử & Đóng gói: ", "Kiểm thử các luồng chức năng, sửa lỗi phát sinh, hoàn thiện tài liệu báo cáo Đồ án cơ sở 2 và slide bảo vệ.")
    ]
    for label, desc in tasks_content:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        p_b.paragraph_format.line_spacing = 1.18
        p_b.paragraph_format.left_indent = Pt(12)
        r = p_b.add_run("● ")
        set_font(r, "Times New Roman", size_pt=8.0, bold=False)
        r = p_b.add_run(label)
        set_font(r, "Times New Roman", size_pt=9.5, bold=True)
        r = p_b.add_run(desc)
        set_font(r, "Times New Roman", size_pt=9.5, bold=False)
    
    # ==========================================
    # NGẮT TRANG SANG TRANG 2 CHO PHẦN KẾ HOẠCH & CHỮ KÝ
    # ==========================================
    doc.add_page_break()
    
    # ==========================================
    # 5. KẾ HOẠCH THỰC HIỆN (Tiến độ 6 tuần)
    # ==========================================
    p_sec5 = doc.add_paragraph()
    p_sec5.paragraph_format.space_before = Pt(0)
    p_sec5.paragraph_format.space_after = Pt(3)
    p_sec5.paragraph_format.line_spacing = 1.2
    r = p_sec5.add_run("5. Kế hoạch thực hiện: ")
    set_font(r, "Times New Roman", size_pt=11.0, bold=True)
    r = p_sec5.add_run("(Tiến độ 6 tuần - Phân công nhóm 02 sinh viên)")
    set_font(r, "Times New Roman", size_pt=9.5, bold=False, italic=True)
    
    # Table 2: Bảng kế hoạch 6 tuần
    # Tỷ lệ cột: Cột 1: 18% (84pt), Cột 2: 60% (280pt), Cột 3: 22% (103pt)
    table2 = doc.add_table(rows=7, cols=3)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table2, border_color="94A3B8", border_sz=4, inside=True)
    
    col_widths = [Pt(84), Pt(280), Pt(103)]
    for row in table2.rows:
        for j, w in enumerate(col_widths):
            row.cells[j].width = w
            set_cell_margins(row.cells[j], top=35, bottom=35, left=60, right=60)
    
    # Header Table 2: Nền xanh đậm 1E3A8A, chữ trắng in đậm
    headers = ["Thời gian", "Nội dung thực hiện", "Người phụ trách chính"]
    for j, h_text in enumerate(headers):
        cell = table2.rows[0].cells[j]
        set_cell_shading(cell, "1E3A8A")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(h_text)
        set_font(r, "Times New Roman", size_pt=9.5, bold=True, color_rgb=RGBColor(0xFF, 0xFF, 0xFF))
    
    weeks_data = [
        (
            "Tuần 1\n(12/10/2026 -\n18/10/2026)",
            "● Khảo sát yêu cầu nghiệp vụ mạng xã hội; phân tích các chức năng chính.\n"
            "● Thiết kế sơ đồ quan hệ thực thể (ERD) và cài đặt CSDL PostgreSQL qua Prisma.\n"
            "● Khởi tạo cấu trúc dự án Frontend (Next.js) và Backend (NestJS).",
            "Hà Văn Phong\n&\n<Thành viên 2>\n(Phối hợp chung)"
        ),
        (
            "Tuần 2\n(19/10/2026 -\n25/10/2026)",
            "● Backend: Xây dựng module Xác thực (Đăng ký, Đăng nhập, mã hóa mật khẩu, JWT) và gửi OTP email.\n"
            "● Frontend: Thiết kế giao diện Đăng ký, Đăng nhập và Modal xác thực mã OTP.",
            "● Hà Văn Phong\n(Backend Auth)\n● <Thành viên 2>\n(Frontend Auth)"
        ),
        (
            "Tuần 3\n(26/10/2026 -\n01/11/2026)",
            "● Backend: Xây dựng module Người dùng (Hồ sơ cá nhân, đổi mật khẩu), tích hợp upload ảnh avatar lên Cloudinary.\n"
            "● Frontend: Thiết kế trang Trang cá nhân (Profile), chỉnh sửa thông tin và upload ảnh đại diện.",
            "● Hà Văn Phong\n(Backend Profile)\n● <Thành viên 2>\n(Frontend Profile)"
        ),
        (
            "Tuần 4\n(02/11/2026 -\n08/11/2026)",
            "● Backend: Xây dựng module Bài viết (Đăng bài viết mới, tải ảnh/video, xem danh sách bài viết, sửa và xóa bài).\n"
            "● Frontend: Thiết kế form Soạn bài viết có đính kèm ảnh và hiển thị danh sách bài viết trên Bảng tin (Feed).",
            "● Hà Văn Phong\n(Backend Bài viết)\n● <Thành viên 2>\n(Frontend Bài viết)"
        ),
        (
            "Tuần 5\n(09/11/2026 -\n15/11/2026)",
            "● Backend: Xây dựng module Tương tác (Thích bài viết, Bình luận bài viết và Theo dõi người dùng khác).\n"
            "● Frontend: Thiết kế khung Bình luận dưới bài viết, nút Thả tim và danh sách người theo dõi.",
            "● Hà Văn Phong\n(Backend Tương tác)\n● <Thành viên 2>\n(Frontend Tương tác)"
        ),
        (
            "Tuần 6\n(16/11/2026 -\n22/11/2026)",
            "● Kiểm thử toàn bộ các chức năng của hệ thống, sửa lỗi và tối ưu giao diện trên thiết bị di động.\n"
            "● Đóng gói mã nguồn, hoàn thiện báo cáo Đề cương Đồ án cơ sở 2 và chuẩn bị slide bảo vệ đề tài.",
            "Hà Văn Phong\n&\n<Thành viên 2>\n(Cả 2 thành viên)"
        )
    ]
    
    for row_idx, (w_time, w_content, w_owner) in enumerate(weeks_data, start=1):
        row = table2.rows[row_idx]
        
        # Col 0: Thời gian
        c0 = row.cells[0]
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing = 1.1
        lines = w_time.split("\n")
        r = p0.add_run(lines[0] + "\n")
        set_font(r, "Times New Roman", size_pt=9.0, bold=True)
        r = p0.add_run("\n".join(lines[1:]))
        set_font(r, "Times New Roman", size_pt=8.5, bold=False, italic=True)
        
        # Col 1: Nội dung thực hiện
        c1 = row.cells[1]
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.15
        r = p1.add_run(w_content)
        set_font(r, "Times New Roman", size_pt=9.0, bold=False)
        
        # Col 2: Người phụ trách chính
        c2 = row.cells[2]
        c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        r = p2.add_run(w_owner)
        set_font(r, "Times New Roman", size_pt=8.5, bold=False)
    
    # ==========================================
    # PHẦN KÝ TÊN: Table 3 (Không viền, 2 cột)
    # ==========================================
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(4)
    p_spacer.paragraph_format.space_after = Pt(2)
    
    table3 = doc.add_table(rows=1, cols=2)
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table3, border_color=None)
    
    col_w_sign = Pt(content_width_pt / 2.0)
    for row in table3.rows:
        row.cells[0].width = col_w_sign
        row.cells[1].width = col_w_sign
        set_cell_margins(row.cells[0], top=10, bottom=10, left=30, right=30)
        set_cell_margins(row.cells[1], top=10, bottom=10, left=30, right=30)
    
    # Cột trái Table 3: Giảng viên hướng dẫn
    cell_gv = table3.rows[0].cells[0]
    p_gv = cell_gv.paragraphs[0]
    p_gv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_gv.paragraph_format.line_spacing = 1.15
    p_gv.paragraph_format.space_after = Pt(0)
    
    r = p_gv.add_run("Ngày ..... tháng ..... năm 2026\n")
    set_font(r, "Times New Roman", size_pt=9.5, bold=False, italic=True)
    r = p_gv.add_run("GIẢNG VIÊN HƯỚNG DẪN\n")
    set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p_gv.add_run("(ký và ghi rõ họ tên)\n\n\n\n")
    set_font(r, "Times New Roman", size_pt=9.0, bold=False, italic=True)
    r = p_gv.add_run("TS. Nguyễn Đức Hiển")
    set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    
    # Cột phải Table 3: 2 Sinh viên thực hiện
    cell_sv = table3.rows[0].cells[1]
    p_sv = cell_sv.paragraphs[0]
    p_sv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sv.paragraph_format.line_spacing = 1.15
    p_sv.paragraph_format.space_after = Pt(0)
    
    r = p_sv.add_run("Đà Nẵng, ngày 08 tháng 10 năm 2026\n")
    set_font(r, "Times New Roman", size_pt=9.5, bold=False, italic=True)
    r = p_sv.add_run("SINH VIÊN THỰC HIỆN\n")
    set_font(r, "Times New Roman", size_pt=10.0, bold=True)
    r = p_sv.add_run("(ký và ghi rõ họ tên)\n\n\n\n")
    set_font(r, "Times New Roman", size_pt=9.0, bold=False, italic=True)
    
    r = p_sv.add_run("Hà Văn Phong           <Họ và tên SV 2>\n")
    set_font(r, "Times New Roman", size_pt=9.5, bold=True)
    r = p_sv.add_run("(Trưởng nhóm)                (Thành viên)")
    set_font(r, "Times New Roman", size_pt=8.5, bold=False, italic=True)
    
    return doc

def reorder_zip_entry(docx_path):
    with zipfile.ZipFile(docx_path, 'r') as zin:
        names = zin.namelist()
        ordered_names = [n for n in names if n == '[Content_Types].xml'] + [n for n in names if n != '[Content_Types].xml']
        data = {n: zin.read(n) for n in names}
    
    with zipfile.ZipFile(docx_path, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
        for n in ordered_names:
            zout.writestr(n, data[n])

def main():
    base_dir = "/home/phongha/Code/mxh"
    out_file1 = os.path.join(base_dir, "de_cuong_do_an_co_so_2_mxh.docx")
    out_file2 = os.path.join(base_dir, "25ITB126_HaVanPhong_DeCuong_DoAnCoSo2.docx")
    
    print("1. Đang khởi tạo tài liệu...")
    doc = create_document()
    
    print(f"2. Đang lưu file gốc: {out_file1}")
    doc.save(out_file1)
    reorder_zip_entry(out_file1)
    
    print(f"3. Đang sao chép file định danh: {out_file2}")
    shutil.copyfile(out_file1, out_file2)
    reorder_zip_entry(out_file2)
    
    print("4. Đang chuyển đổi sang PDF qua LibreOffice headless...")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", out_file1, "--outdir", base_dir], check=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", out_file2, "--outdir", base_dir], check=True)
    
    print("Hoàn tất tạo file docx và pdf thành công!")

if __name__ == "__main__":
    main()
